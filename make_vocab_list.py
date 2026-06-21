from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ── vocabulary lists ──────────────────────────────────────────

ann_col = [
    ('__h__', 'ANN — Artificial Neural Networks'),
    ('__s__', 'Core Concepts'),
    'Artificial Neural Network (ANN)',
    'Neuron / Node',
    'Weight',
    'Bias',
    'Input Layer',
    'Hidden Layer',
    'Output Layer',
    ('__s__', 'Activation Functions'),
    'Activation Function',
    'Sigmoid',
    'ReLU (Rectified Linear Unit)',
    'Tanh (Hyperbolic Tangent)',
    'Softmax',
    ('__s__', 'Training'),
    'Forward Pass',
    'Loss Function',
    'Binary Cross-Entropy',
    'Backpropagation',
    'Gradient Descent',
    'Learning Rate (η)',
    'Epoch',
    'Batch Size',
    'Normalisation',
    ('__s__', 'Image Processing'),
    'Grayscale Image',
    'RGB Image',
    'Flatten',
    'Weighted Sum (Σ)',
]

cnn_col = [
    ('__h__', 'CNN — Convolutional Neural Networks'),
    ('__s__', 'Dataset & Input'),
    'Dataset',
    'Pixel Value',
    'Preprocessing',
    ('__s__', 'Convolutional Layers'),
    'Convolutional Layer',
    'Filter / Kernel',
    'Feature Map (Activation Map)',
    'Stride',
    'Padding',
    'Weight Sharing',
    ('__s__', 'Activation & Pooling'),
    'ReLU (Rectified Linear Unit)',
    'Max Pooling',
    'Receptive Field',
    ('__s__', 'Dense Layers & Output'),
    'Flatten Layer',
    'Dense Layer (Fully Connected)',
    'Dropout',
    'Overfitting',
    'Softmax Output',
    ('__s__', 'Training'),
    'Loss Function',
    'Backpropagation',
    'Batch Normalisation',
    'Training Loop',
]

yolo_col = [
    ('__h__', 'YOLO — Object Detection'),
    ('__s__', 'YOLO Fundamentals'),
    'YOLO (You Only Look Once)',
    'Single-Pass Detection',
    'R-CNN (Region-based CNN)',
    'Darknet Framework',
    'Anchor Boxes',
    'Grid System',
    'mAP (mean Average Precision)',
    'FPS (Frames Per Second)',
    ('__s__', 'YOLOv8 Architecture'),
    'Backbone',
    'Neck (PAN-FPN)',
    'Head',
    'CSPDarknet',
    'C2f Module',
    'SPPF',
    'SiLU Activation',
    'Decoupled Head',
    ('__s__', 'Detection & Evaluation'),
    'Bounding Box',
    'Confidence Score',
    'IoU (Intersection over Union)',
    'NMS (Non-Maximum Suppression)',
    'Anchor-Free Detection',
    'DFL Loss',
    'TAL (Task-Aligned Learning)',
]

cols = [ann_col, cnn_col, yolo_col]
accents = ['1f6b35', '1a5a9a', 'b55a10']   # dark green, dark blue, dark orange

# ── helpers ────────────────────────────────────────────────────

def set_cell_no_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:val'), 'none')
        node.set(qn('w:sz'), '0')
        node.set(qn('w:space'), '0')
        node.set(qn('w:color'), 'auto')
        tcBorders.append(node)
    tcPr.append(tcBorders)

def set_col_width(cell, width_cm):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(int(width_cm * 567)))   # 567 twips per cm
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)

# ── build doc ─────────────────────────────────────────────────

doc = Document()

# margins: very narrow to maximise space
for sec in doc.sections:
    sec.top_margin    = Cm(1.2)
    sec.bottom_margin = Cm(1.2)
    sec.left_margin   = Cm(1.4)
    sec.right_margin  = Cm(1.4)
    sec.page_width    = Cm(21.0)   # A4
    sec.page_height   = Cm(29.7)

# page title
pt = doc.add_paragraph()
pt.alignment = WD_ALIGN_PARAGRAPH.CENTER
pt.paragraph_format.space_after = Pt(6)
rt = pt.add_run('Vocabulary Quick Reference  —  ANN · CNN · YOLO')
rt.bold = True
rt.font.size = Pt(12)
rt.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

# thin rule
rule = doc.add_paragraph()
rule.paragraph_format.space_before = Pt(0)
rule.paragraph_format.space_after  = Pt(6)
rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = rule.add_run('─' * 110)
rr.font.size = Pt(5)
rr.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

# calculate row count = max column length
max_rows = max(len(c) for c in cols)

tbl = doc.add_table(rows=max_rows, cols=3)
tbl.style = 'Table Grid'

# remove all table borders globally
tbl_pr = tbl._tbl.tblPr
tbl_borders = OxmlElement('w:tblBorders')
for side in ('top','left','bottom','right','insideH','insideV'):
    node = OxmlElement(f'w:{side}')
    node.set(qn('w:val'), 'none')
    node.set(qn('w:sz'), '0')
    node.set(qn('w:color'), 'auto')
    tbl_borders.append(node)
tbl_pr.append(tbl_borders)

COL_W = 5.9   # cm per column (3 × 5.9 = 17.7 cm fits in 18.2 cm printable width)

for col_idx, (items, hex_accent) in enumerate(zip(cols, accents)):
    r, g, b = bytes.fromhex(hex_accent)
    accent_rgb = RGBColor(r, g, b)

    for row_idx, item in enumerate(items):
        cell = tbl.cell(row_idx, col_idx)
        set_cell_no_border(cell)
        set_col_width(cell, COL_W)

        p = cell.paragraphs[0]
        p.clear()

        if isinstance(item, tuple) and item[0] == '__h__':
            # Column header
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(4)
            run = p.add_run(item[1])
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = accent_rgb

        elif isinstance(item, tuple) and item[0] == '__s__':
            # Section sub-header
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after  = Pt(1)
            run = p.add_run(f'  {item[1]}')
            run.bold = True
            run.font.size = Pt(7.5)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        else:
            # Term
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after  = Pt(0)
            p.paragraph_format.left_indent  = Pt(10)
            run = p.add_run(f'• {item}')
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # fill empty cells below
    for row_idx in range(len(items), max_rows):
        cell = tbl.cell(row_idx, col_idx)
        set_cell_no_border(cell)
        set_col_width(cell, COL_W)
        cell.paragraphs[0].clear()

out = '/home/user/git_test/vocabulary_quick_reference.docx'
doc.save(out)
print(f'Saved {out}')
